#!/usr/bin/env python3
"""STEP 11 -- per-cell-type depth in the Qiu et al. 2024 developmental graph, used ONLY as a
guidepost for which member of a pair to list first.

Depth = shortest-path length from "Oocyte" in the prior developmental graph (Suppl. Tables 20 + 22,
same-name nodes merged across subsystems), so a smaller depth means the cell type appears earlier in
that reference. Name matching is the same normalisation used by
draft_hierarchy/hierarchy/code/hv6_07_prior.py.

This does NOT feed any statistic. It only fixes the order in which the two cell types of a pair are
written, so that the earlier one comes first. Ties and unmapped types fall back to alphabetical
order, which keeps the output deterministic.

Writes out/qiu_celltype_depth.csv.
"""
import openpyxl, warnings, re, os, csv, numpy as np, networkx as nx

import os as _os
_REPO = _os.path.abspath(_os.path.join(_os.path.dirname(_os.path.abspath(__file__)), "..", ".."))

def _support(name, env=None, hint=None):
    """Resolve an input under support_data/, overridable by env var."""
    if env:
        v = _os.environ.get(env)
        if v:
            return v
    p = _os.path.join(_REPO, "support_data", name)
    if not _os.path.exists(p) and hint:
        raise FileNotFoundError(f"{p} not found. {hint}")
    return p

warnings.filterwarnings("ignore")

HERE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
XLSX = _os.environ.get("QIU2024_SUPP_XLSX", _os.path.join(_REPO, "support_data", "external", "41586_2024_7069_MOESM4_ESM.xlsx"))

wb = openpyxl.load_workbook(XLSX, read_only=True)
name_of = {}
for r in list(wb["Table.20"].iter_rows(values_only=True))[3:]:
    if r[1] is not None: name_of[str(r[1]).strip()] = str(r[2]).strip()
edges = [(str(r[1]).strip(), str(r[2]).strip())
         for r in list(wb["Table.22"].iter_rows(values_only=True))[3:] if r[1] and r[2]]
G = nx.Graph()
for a, b in edges:
    na, nb = name_of.get(a, a), name_of.get(b, b)
    if na != nb: G.add_edge(na, nb)
depth = nx.single_source_shortest_path_length(G, "Oocyte")

def norm(s):
    s = s.lower().strip().replace("progenitors", "prog").replace("progenitor", "prog")
    s = re.sub(r"[()+\-–,]", " ", s).replace("cells", "").replace("cell", "")
    s = re.sub(r"\bprog\w*\b", "prog", s); s = re.sub(r"\s+", " ", s).strip()
    return s
prior_norm = {}
for nm in set(name_of.values()):
    prior_norm.setdefault(norm(nm), nm)
MANUAL = {"GABAergic neurons": "GABAergic neurons (after E13.0)",
          "Glutamatergic neurons": "Glutamatergic neurons (after E13.0)",
          "Spinal cord dorsal progenitors": "Spinal cord dorsal progenitors (after E13.0)",
          "Lateral plate and intermediate mesoderm": "Lateral plate mesoderm"}

types = list(np.load(f"{HERE}/out/c7_timesweep.npz", allow_pickle=True)["types"])
rows, nmap = [], 0
for ct in types:
    pn = MANUAL.get(ct) or prior_norm.get(norm(ct))
    dep = depth.get(pn) if pn in G else None
    nmap += dep is not None
    rows.append((ct, pn or "", "" if dep is None else dep))
with open(f"{HERE}/out/qiu_celltype_depth.csv", "w", newline="") as fh:
    w = csv.writer(fh); w.writerow(["celltype", "qiu_node", "qiu_depth_from_oocyte"])
    for r in sorted(rows, key=lambda r: (r[2] == "", r[2], r[0])): w.writerow(r)
print(f"mapped {nmap}/{len(types)} cell types into the prior graph")
print(f"depth range: {min(d for _,_,d in rows if d!='')} .. {max(d for _,_,d in rows if d!='')}")
print("\nunmapped (will fall back to alphabetical ordering):")
for ct, pn, dep in rows:
    if dep == "": print(f"   {ct}")
print("\nearliest 12 by prior-graph depth:")
for ct, pn, dep in sorted([r for r in rows if r[2] != ""], key=lambda r: r[2])[:12]:
    print(f"   depth {dep:>2}  {ct:<42} -> {pn}")
print(f"\nwrote out/qiu_celltype_depth.csv")

#!/usr/bin/env python3
"""Q1: validate the imputed ancestor->descendant transitions against the MANUALLY CURATED
Qiu et al. 2024 developmental graph (Suppl. Tables 20 + 22).

Why this is a fair test: the imputation used only the atlas *expression* data (joint PCs +
MNN). The curated graph -- literature-reviewed edges between cell states -- was never an
input. So graph adjacency is an independent ground truth for the inferred successions.

Graph construction is a faithful reuse of clade_cooccur_final/code/c9_qiu.py (same tables,
same same-name merge, same norm()/MANUAL name matching) so numbers are comparable to the
Mantel rho=0.47-0.51 reported there.
"""
import openpyxl, warnings, re, csv, sys, numpy as np, networkx as nx

import os as _os
_REPO = _os.path.abspath(_os.path.join(_os.path.dirname(_os.path.abspath(__file__)), "..", ".."))

def _support(name, env=None, hint=None):
    """Resolve an input under support_data/, overridable by env var."""
    if env:
        v = _os.environ.get(env)
        if v:
            return v
    p = _os.path.join(_REPO, "support_data", name)
    if not _os.path.exists(p) and hint:
        raise FileNotFoundError(f"{p} not found. {hint}")
    return p

warnings.filterwarnings("ignore")
XLSX=_os.environ.get("QIU2024_SUPP_XLSX", _os.path.join(_REPO, "support_data", "external", "41586_2024_7069_MOESM4_ESM.xlsx"))
RNG=np.random.default_rng(0)

# ---------------- Qiu curated graph (undirected G + directed DG) ----------------
wb=openpyxl.load_workbook(XLSX,read_only=True)
name_of={}
for r in list(wb["Table.20"].iter_rows(values_only=True))[3:]:
    if r[1] is not None: name_of[str(r[1]).strip()]=str(r[2]).strip()
raw=[(str(r[1]).strip(),str(r[2]).strip()) for r in list(wb["Table.22"].iter_rows(values_only=True))[3:] if r[1] and r[2]]
G=nx.Graph(); DG=nx.DiGraph()
for a,b in raw:
    na,nb=name_of.get(a,a),name_of.get(b,b)
    if na!=nb: G.add_edge(na,nb); DG.add_edge(na,nb)
depth=nx.single_source_shortest_path_length(G,"Oocyte")
print(f"Qiu curated graph: {G.number_of_nodes()} nodes, {G.number_of_edges()} undirected edges, "
      f"{DG.number_of_edges()} directed; connected={nx.is_connected(G)}")

def norm(s):
    s=s.lower().strip().replace("progenitors","prog").replace("progenitor","prog")
    s=re.sub(r"[()+\-–,]"," ",s).replace("cells","").replace("cell","")
    s=re.sub(r"\bprog\w*\b","prog",s); s=re.sub(r"\s+"," ",s).strip(); return s
prior_norm={}
for nm in set(name_of.values()): prior_norm.setdefault(norm(nm),nm)
MANUAL={"GABAergic neurons":"GABAergic neurons (after E13.0)",
        "Glutamatergic neurons":"Glutamatergic neurons (after E13.0)",
        "Spinal cord dorsal progenitors":"Spinal cord dorsal progenitors (after E13.0)",
        "Lateral plate and intermediate mesoderm":"Lateral plate mesoderm"}
def pmap(ct): return MANUAL.get(ct) or prior_norm.get(norm(ct))

# ---------------- our inferred edges ----------------
def load(fn):
    rows=list(csv.DictReader(open(fn)))
    return [(r["anc"],r["desc"],float(r["n_fwd"]),float(r["asym"]),float(r["d_gap"]),
             r["keep"]=="True") for r in rows]
ALL=load("out_edges_all.csv"); FINAL=load("out_edges_final.csv")
types=sorted({t for e in ALL for t in e[:2]})
mapped={t:pmap(t) for t in types}
inG=[t for t in types if mapped[t] in G]
print(f"our cell types in the edge tables: {len(types)}; mapped into the curated graph: {len(inG)}")
print(f"  unmapped: {sorted(t for t in types if mapped[t] not in G)[:12]}")

sp=dict(nx.all_pairs_shortest_path_length(G))
def qdist(a,b):
    pa,pb=mapped.get(a),mapped.get(b)
    if pa not in G or pb not in G: return None
    return sp[pa].get(pb)

# ---------------- background: all ordered pairs of mappable types ----------------
bg=[qdist(a,b) for a in inG for b in inG if a!=b]
bg=[d for d in bg if d is not None]
def summ(ds,lab):
    ds=[d for d in ds if d is not None]; n=len(ds); a=np.array(ds)
    return (f"{lab:38} n={n:5}  adjacent(d=1) {100*(a==1).mean():5.1f}%  "
            f"d<=2 {100*(a<=2).mean():5.1f}%  d<=3 {100*(a<=3).mean():5.1f}%  median d={np.median(a):.1f}")
print("\n=== Qiu-graph distance between the two endpoints of an inferred transition ===")
print(summ(bg,"BACKGROUND (all mappable pairs)"))
for lab,E in [("all 434 candidate transitions",ALL),("66 filtered high-confidence edges",FINAL)]:
    print(summ([qdist(a,b) for a,b,*_ in E],lab))
# stratify the candidate set by our own confidence stats
for thr in (0.2,0.4,0.6):
    sel=[e for e in ALL if abs(e[3])>=thr]
    print(summ([qdist(a,b) for a,b,*_ in sel],f"  candidates with |asym| >= {thr}"))
for gap in (0.5,1.0):
    sel=[e for e in ALL if e[4]>=gap]
    print(summ([qdist(a,b) for a,b,*_ in sel],f"  candidates with d_gap >= {gap} d"))

# ---------------- significance: label-permutation Mantel-style test ----------------
print("\n=== significance (permute cell-type labels on the curated graph, 9,999 draws) ===")
pos={t:mapped[t] for t in inG}
names=[mapped[t] for t in inG]
def frac_adj(E,perm=None):
    m=dict(zip(inG,names)) if perm is None else dict(zip(inG,perm))
    ds=[]
    for a,b,*_ in E:
        pa,pb=m.get(a),m.get(b)
        if pa in G and pb in G: ds.append(sp[pa].get(pb))
    ds=np.array([d for d in ds if d is not None]); return (ds==1).mean(), np.median(ds)
for lab,E in [("all 434 candidates",ALL),("66 filtered edges",FINAL)]:
    obs,obsmed=frac_adj(E)
    null=np.empty(9999); nullmed=np.empty(9999)
    for i in range(9999):
        p=list(names); RNG.shuffle(p); null[i],nullmed[i]=frac_adj(E,p)
    pv=(1+(null>=obs).sum())/(1+len(null))
    print(f"  {lab:22} adjacent {100*obs:5.1f}%  null {100*null.mean():4.1f}% +/- {100*null.std():.1f}  "
          f"enrichment {obs/null.mean():5.1f}x  p={pv:.5f}   median d {obsmed:.1f} vs null {np.median(nullmed):.1f}")

# ---------------- does our inferred DIRECTION agree with the curated graph? ----------------
print("\n=== direction agreement, for inferred edges that ARE curated-graph adjacent ===")
print("    (curated direction = away from Oocyte, i.e. increasing graph depth)")
for lab,E in [("all candidates",ALL),("66 filtered edges",FINAL)]:
    ok=bad=tie=0
    for a,b,*_ in E:
        pa,pb=mapped.get(a),mapped.get(b)
        if pa not in G or pb not in G or sp[pa].get(pb)!=1: continue
        da,db=depth.get(pa),depth.get(pb)
        if da is None or db is None or da==db: tie+=1
        elif da<db: ok+=1
        else: bad+=1
    n=ok+bad
    print(f"  {lab:18} anc shallower than desc: {ok}/{n} = {100*ok/max(n,1):.1f}%  "
          f"(reversed {bad}, equal-depth {tie})")

# ---------------- per-edge detail for the 66 ----------------
print("\n=== the 66 filtered edges vs the curated graph ===")
print(f"  {'d':>2}  {'n_fwd':>7} {'asym':>5}  edge")
rows=sorted(((qdist(a,b),n,asym,a,b) for a,b,n,asym,*_ in FINAL),
            key=lambda r:(99 if r[0] is None else r[0],-r[1]))
for d,n,asym,a,b in rows:
    print(f"  {('-' if d is None else d):>2}  {int(n):7,} {asym:5.2f}  {a} -> {b}")

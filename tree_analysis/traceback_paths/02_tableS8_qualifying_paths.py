#!/usr/bin/env python3
"""Q7: generate Supplementary Table S7 -- the qualifying ancestor-to-descendant paths.

Definition (identical to the manuscript and to q6_qualifying_site.py with USE_BINOM = False):
a traceback path qualifies if
  (i)   it is heterotypic (>= 2 distinct states) and non-recursive (no state appears twice);
  (ii)  it is carried by >= 1% of that cell type's traced cells and by >= 5 cells; and
  (iii) every step is directionally asymmetric in the tree: >= 50 observed ancestor->descendant
        transitions between the two states, with asymmetry (f-r)/(f+r) >= 0.2.
All three use only the phylogeny and the imputed labels. The curated Qiu et al. 2024 graph and
the germ-layer map are reported as independent annotation, never used to select paths.

Writes TableS7_qualifying_paths.csv and .xlsx (Legend / Summary / Paths sheets).
"""
import numpy as np, csv, collections, re, warnings, openpyxl, networkx as nx, os
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter
warnings.filterwarnings("ignore")

BASE = "/Users/jay.shendure/Dropbox/claude/current/final_push/flowsite_v8/"
NPZ = "/Users/jay.shendure/Dropbox/claude/top_to_bottom_phylogeny/out/mergedtree_dttpq_v8.npz"
TXT = BASE + "pd_nodes_infer_200.txt"
GL = ("/Users/jay.shendure/Dropbox/claude/penultimate_tree_build/heterotypic_siblings/"
      "data/germ_layer_map_validated.csv")
XLSX_IN = "/Users/jay.shendure/Dropbox/claude/penultimate_clade_k_analysis/41586_2024_7069_MOESM4_ESM (4).xlsx"
OUT_CSV = BASE + "TableS7_qualifying_paths.csv"
OUT_XLSX = BASE + "TableS7_qualifying_paths.xlsx"
SHARE, MINCELLS, ASYM, MINTOT = 0.01, 5, 0.2, 50

# ---------------------------------------------------------------- curated graph
wb = openpyxl.load_workbook(XLSX_IN, read_only=True); name_of = {}
for r in list(wb["Table.20"].iter_rows(values_only=True))[3:]:
    if r[1] is not None: name_of[str(r[1]).strip()] = str(r[2]).strip()
G = nx.Graph()
for r in list(wb["Table.22"].iter_rows(values_only=True))[3:]:
    if r[1] and r[2]:
        a = name_of.get(str(r[1]).strip(), str(r[1]).strip())
        b = name_of.get(str(r[2]).strip(), str(r[2]).strip())
        if a != b: G.add_edge(a, b)
depth = nx.single_source_shortest_path_length(G, "Oocyte")
def norm(s):
    s = s.lower().strip().replace("progenitors", "prog").replace("progenitor", "prog")
    s = re.sub(r"[()+\-–,]", " ", s).replace("cells", "").replace("cell", "")
    s = re.sub(r"\bprog\w*\b", "prog", s); return re.sub(r"\s+", " ", s).strip()
pn = {}
for nm in set(name_of.values()): pn.setdefault(norm(nm), nm)
MAN = {"GABAergic neurons": "GABAergic neurons (after E13.0)",
       "Glutamatergic neurons": "Glutamatergic neurons (after E13.0)",
       "Spinal cord dorsal progenitors": "Spinal cord dorsal progenitors (after E13.0)",
       "Lateral plate and intermediate mesoderm": "Lateral plate mesoderm"}
def pm(c): return MAN.get(c) or pn.get(norm(c), c)
sp = dict(nx.all_pairs_shortest_path_length(G))
print(f"curated graph: {G.number_of_nodes()} nodes, {G.number_of_edges()} undirected edges")

# ---------------------------------------------------------------- tree + calls
z = np.load(NPZ, allow_pickle=True)
parent = z["parent"].astype(np.int64); is_leaf = z["is_leaf"]
Nflat = parent.size; Ntip = int(is_leaf.sum()); Nint = Nflat - Ntip
ape = np.empty(Nflat, dtype=np.int64)
ape[is_leaf] = np.arange(1, Ntip + 1); ape[~is_leaf] = Ntip + np.arange(1, Nint + 1)
lab_a = np.empty(Nflat + 1, dtype=object)
with open(TXT) as f:
    r = csv.reader(f, delimiter="\t"); next(r)
    for nid, day, ct, cid in r: lab_a[int(nid.split("_")[1])] = ct
lab = lab_a[ape]
size = np.ones(Nflat, dtype=np.int64)
for v in range(Nflat - 1, 0, -1): size[parent[v]] += size[v]
root = int(np.where(parent < 0)[0][0])
kids = collections.defaultdict(list)
for v in range(Nflat):
    if parent[v] >= 0: kids[parent[v]].append(v)
side = np.full(Nflat, -1, dtype=np.int8)
for j, c in enumerate(sorted(kids[root], key=lambda c: -size[c])[:2]):
    st = [c]
    while st:
        u = st.pop(); side[u] = j; st.extend(kids[u])

par = parent.tolist(); L = lab.tolist(); leaf = is_leaf.tolist()
paths = collections.Counter(); denom = collections.Counter()
origins = collections.defaultdict(set); sides = collections.defaultdict(set)
T2 = collections.Counter(); nca = np.full(Nflat, -1, dtype=np.int64)
for v in range(Nflat):
    p = par[v]
    if p < 0: continue
    nca[v] = p if (L[p] is not None and L[p] != "missing") else nca[p]
    a = nca[v]
    if a >= 0 and L[v] is not None and L[v] != "missing" and L[a] != L[v]:
        T2[(L[a], L[v])] += 1
for tip in range(Nflat):
    if not leaf[tip]: continue
    T = L[tip]
    if T is None or T == "missing": continue
    denom[T] += 1
    seq = [T]; nodes = [tip]; v = par[tip]
    while v >= 0:
        s = L[v]
        if s is not None and s != "missing" and s != seq[-1]: seq.append(s); nodes.append(v)
        v = par[v]
    if len(seq) < 2 or len(set(seq)) != len(seq): continue
    p = tuple(reversed(seq))
    paths[p] += 1; origins[p].add(nodes[-1]); sides[p].add(int(side[tip]))
byT = collections.defaultdict(list)
for p, n in paths.items(): byT[p[-1]].append((n, p))
traced = {T: sum(n for n, _ in lst) for T, lst in byT.items()}
gm = {}
for row in csv.DictReader(open(GL)):
    if row.get("germ_layer"): gm[row["celltype"]] = row["germ_layer"]

def asym(a, b):
    f, r = T2[(a, b)], T2[(b, a)]
    return (None, f, r) if f + r < MINTOT else ((f - r) / (f + r), f, r)

# ---------------------------------------------------------------- qualify
QUAL = {}
for T, lst in byT.items():
    keep = []
    for n, p in sorted(lst, key=lambda x: -x[0]):
        if n < MINCELLS or n / traced[T] < SHARE: continue
        A = [asym(p[j], p[j + 1])[0] for j in range(len(p) - 1)]
        if all(a is not None and a >= ASYM for a in A): keep.append((n, p))
    if keep: QUAL[T] = keep
NP = sum(len(v) for v in QUAL.values()); NC = sum(n for v in QUAL.values() for n, _ in v)
print(f"qualifying: {NP} paths, {len(QUAL)} cell types, {NC:,} cells")

# ---------------------------------------------------------------- rows
def qd(a, b):
    pa, pb = pm(a), pm(b)
    if pa not in G or pb not in G: return None
    return sp[pa].get(pb)
def qdirn(a, b):
    if qd(a, b) != 1: return None
    da, db = depth.get(pm(a)), depth.get(pm(b))
    if da is None or db is None or da == db: return None
    return da < db

rows = []
for T, v in QUAL.items():
    for n, p in v:
        st = list(zip(p[:-1], p[1:]))
        A = [asym(a, b) for a, b in st]
        ds = [qd(a, b) for a, b in st]
        dn = [qdirn(a, b) for a, b in st]
        orient = [x for x in dn if x is not None]
        gl = [gm.get(s, "") for s in p]
        rows.append({
            "celltype": T, "germ_layer": gm.get(T, "unassigned"),
            "n_cells_celltype": denom[T], "n_cells_traced": traced[T],
            "n_routes_celltype": len(byT[T]),
            "path": " -> ".join(p), "n_states": len(p), "n_steps": len(st),
            "n_cells_path": n,
            "pct_of_traced": round(100 * n / traced[T], 2),
            "pct_of_celltype": round(100 * n / denom[T], 2),
            "n_independent_origins": len(origins[p]),
            "both_blastomeres": "yes" if len(sides[p]) == 2 else "no",
            "min_step_asymmetry": round(min(a for a, _, _ in A), 3),
            "step_asymmetry": ";".join(f"{a:+.2f}" for a, _, _ in A),
            "step_transitions_fwd_rev": ";".join(f"{f}/{r}" for _, f, r in A),
            "step_curated_distance": ";".join("NA" if d is None else str(d) for d in ds),
            "all_steps_adjacent": "yes" if (ds and all(d == 1 for d in ds if d is not None)
                                            and any(d is not None for d in ds)) else "no",
            "all_steps_within_two": "yes" if (ds and all(d <= 2 for d in ds if d is not None)
                                              and any(d is not None for d in ds)) else "no",
            "curated_direction_agreement": (f"{sum(orient)}/{len(orient)}" if orient else "NA"),
            "germ_layers": " -> ".join(x if x else "?" for x in gl),
            "single_germ_layer": "yes" if (all(gl) and len(set(gl)) == 1) else
                                 ("no" if all(gl) else "untested"),
        })
GORD = {g: i for i, g in enumerate(
    ["BLOOD", "ENDODERM", "MESODERM", "NEURAL_CREST", "NEUROECTODERM", "SURFACE_ECTO",
     "unassigned"])}
rows.sort(key=lambda r: (GORD.get(r["germ_layer"], 99), -r["n_cells_celltype"],
                         -r["n_cells_path"]))
for i, r in enumerate(rows, 1): r["path_id"] = i
COLS = ["path_id", "celltype", "germ_layer", "n_cells_celltype", "n_cells_traced",
        "n_routes_celltype", "path", "n_states", "n_steps", "n_cells_path", "pct_of_traced",
        "pct_of_celltype", "n_independent_origins", "both_blastomeres", "min_step_asymmetry",
        "step_asymmetry", "step_transitions_fwd_rev", "step_curated_distance",
        "all_steps_adjacent", "all_steps_within_two", "curated_direction_agreement",
        "germ_layers", "single_germ_layer"]
with open(OUT_CSV, "w", newline="") as f:
    w = csv.DictWriter(f, fieldnames=COLS); w.writeheader()
    for r in rows: w.writerow({k: r[k] for k in COLS})
print(f"wrote {OUT_CSV} ({len(rows)} rows)")

# ---------------------------------------------------------------- summary stats
tot_steps = sum(r["n_steps"] for r in rows)
tst = [d for r in rows for d in r["step_curated_distance"].split(";") if d != "NA"]
d1 = sum(1 for d in tst if d == "1"); d2 = sum(1 for d in tst if int(d) <= 2)
allty = sorted({s for r in rows for s in r["path"].split(" -> ")})
mapped = [c for c in allty if pm(c) in G]
bgd = np.array([d for d in (sp[pm(a)].get(pm(b)) for a in mapped for b in mapped if a != b)
                if d is not None])
ok = sum(int(x.split("/")[0]) for r in rows if r["curated_direction_agreement"] != "NA"
         for x in [r["curated_direction_agreement"]])
orn = sum(int(x.split("/")[1]) for r in rows if r["curated_direction_agreement"] != "NA"
          for x in [r["curated_direction_agreement"]])
SUM = [
    ("Qualifying paths", NP), ("Cell types represented", len(QUAL)),
    ("Cell types with any traceable ancestry", len(byT)),
    ("Cells on a qualifying path", NC),
    ("Cells with a traceable, non-recursive ancestry", sum(paths.values())),
    ("Multi-step paths (>= 3 states)", sum(1 for r in rows if r["n_states"] >= 3)),
    ("Paths recovered in both blastomeres", sum(1 for r in rows if r["both_blastomeres"] == "yes")),
    ("Steps in qualifying paths", tot_steps),
    ("Steps testable against the curated graph", len(tst)),
    ("Steps directly adjacent (d=1)", f"{d1} ({100*d1/len(tst):.1f}%)"),
    ("Steps within two curated edges (d<=2)", f"{d2} ({100*d2/len(tst):.1f}%)"),
    ("Background d=1 / d<=2 (labels permuted among these states)",
     f"{100*(bgd==1).mean():.1f}% / {100*(bgd<=2).mean():.1f}%"),
    ("Enrichment d=1 / d<=2", f"{(d1/len(tst))/(bgd==1).mean():.1f}x / "
                              f"{(d2/len(tst))/(bgd<=2).mean():.1f}x"),
    ("Adjacent, orientable steps running in the curated direction", f"{ok}/{orn}"),
    ("Paths entirely within one germ layer",
     f"{sum(1 for r in rows if r['single_germ_layer']=='yes')} of "
     f"{sum(1 for r in rows if r['single_germ_layer'] in ('yes','no'))} testable"),
    ("Curated graph used for scoring",
     f"Qiu et al. 2024 Suppl. Tables 20+22: {G.number_of_nodes()} states, "
     f"{G.number_of_edges()} undirected edges"),
]
for k, v in SUM: print(f"  {k}: {v}")

# ---------------------------------------------------------------- xlsx
LEGEND = [
    ("Supplementary Table S7", ""),
    ("Qualifying ancestor-to-descendant paths inferred from the embryo #3 phylogeny", ""),
    ("", ""),
    ("Definition", "A traceback path qualifies if all three criteria hold. All three use only "
                   "the phylogeny and the imputed cell type labels; no external knowledge of "
                   "developmental biology is used to select paths."),
    ("  (i) form", "Heterotypic (>= 2 distinct states) and non-recursive (no state appears "
                   "twice). Per tip, the lineage is read root-ward keeping only nodes that "
                   "received a call, collapsing contiguous repeats, and ending on the tip's own "
                   "observed E13.5 cell type."),
    ("  (ii) frequency", f"Carried by >= {SHARE:.0%} of that cell type's traced cells and by "
                         f">= {MINCELLS} cells."),
    ("  (iii) directionality", f"Every step (state A -> state B) has >= {MINTOT} observed "
                               f"ancestor-to-descendant transitions between A and B across the "
                               f"whole tree, with asymmetry (n_AB - n_BA)/(n_AB + n_BA) >= "
                               f"{ASYM} (i.e. the observed direction favoured >= 60:40)."),
    ("", ""),
    ("Independent annotation", "The curated graph of Qiu et al. 2024 (Suppl. Tables 20+22) and "
                              "the germ-layer map are reported per path but were never used to "
                              "select paths. Curated distance 1 = the two states are directly "
                              "adjacent; 2 = one intervening state. Because paths keep only "
                              "called nodes and collapse repeats, a single step may legitimately "
                              "span more than one curated edge."),
    ("", ""),
    ("COLUMN", "DESCRIPTION"),
    ("path_id", "Row identifier, 1-based, ordered by germ layer then cell type size."),
    ("celltype", "The observed E13.5 cell type at which the path terminates."),
    ("germ_layer", "Germ layer of the terminal cell type; 'unassigned' where unmapped."),
    ("n_cells_celltype", "Cells with this annotation among the 1,340,794 E13.5 tips."),
    ("n_cells_traced", "Of those, cells with a heterotypic, non-recursive ancestry."),
    ("n_routes_celltype", "Distinct such routes for this cell type, before any frequency or "
                          "directionality filter."),
    ("path", "The path, oldest state to youngest, ending on the observed E13.5 type."),
    ("n_states / n_steps", "States in the path, and transitions between them (n_states - 1)."),
    ("n_cells_path", "Cells following this path."),
    ("pct_of_traced", "n_cells_path as a percentage of n_cells_traced (the criterion (ii) "
                      "denominator)."),
    ("pct_of_celltype", "n_cells_path as a percentage of n_cells_celltype."),
    ("n_independent_origins", "Distinct tree nodes at which this path's earliest state occurs, "
                              "i.e. how many independent lineages support it."),
    ("both_blastomeres", "Whether the path is recovered in both blastomere-derived halves of "
                         "the embryo (independent replication)."),
    ("min_step_asymmetry", "The least asymmetric step in the path; >= 0.2 by construction."),
    ("step_asymmetry", "Per-step asymmetry, semicolon-separated, in path order."),
    ("step_transitions_fwd_rev", "Per-step observed transition counts, forward/reverse."),
    ("step_curated_distance", "Per-step shortest-path distance in the curated graph; NA where "
                              "either state is absent from it."),
    ("all_steps_adjacent", "Every testable step is a direct curated edge (distance 1)."),
    ("all_steps_within_two", "Every testable step is within two curated edges."),
    ("curated_direction_agreement", "Among steps that are curated-adjacent and orientable "
                                    "(the two states sit at different depths from 'Oocyte'), "
                                    "how many run in the curated direction."),
    ("germ_layers", "Germ layer of each state in path order; '?' where unmapped."),
    ("single_germ_layer", "yes / no / untested (untested = at least one state has no germ-layer "
                          "assignment)."),
]
wbo = openpyxl.Workbook()
ws = wbo.active; ws.title = "Legend"
for a, b in LEGEND: ws.append([a, b])
ws.column_dimensions["A"].width = 30; ws.column_dimensions["B"].width = 108
for c in ws["A"]: c.font = Font(bold=True)
for row in ws.iter_rows(min_col=2, max_col=2):
    for c in row: c.alignment = Alignment(wrap_text=True, vertical="top")
ws["A1"].font = Font(bold=True, size=14); ws["A2"].font = Font(italic=True, size=11)

ws2 = wbo.create_sheet("Summary")
ws2.append(["METRIC", "VALUE"])
for k, v in SUM: ws2.append([k, v])
ws2.column_dimensions["A"].width = 62; ws2.column_dimensions["B"].width = 58
for c in ws2[1]: c.font = Font(bold=True)

ws3 = wbo.create_sheet("Paths")
ws3.append(COLS)
for r in rows: ws3.append([r[k] for k in COLS])
hdr = PatternFill("solid", fgColor="EFEFEF")
for c in ws3[1]: c.font = Font(bold=True); c.fill = hdr
ws3.freeze_panes = "A2"
ws3.auto_filter.ref = f"A1:{get_column_letter(len(COLS))}{len(rows)+1}"
widths = {"celltype": 40, "path": 96, "germ_layers": 52, "step_asymmetry": 20,
          "step_transitions_fwd_rev": 24, "step_curated_distance": 20}
for i, k in enumerate(COLS, 1):
    ws3.column_dimensions[get_column_letter(i)].width = widths.get(k, max(11, len(k) + 2))
wbo.save(OUT_XLSX)
print(f"wrote {OUT_XLSX} ({os.path.getsize(OUT_XLSX)/1024:.0f} KB)")
